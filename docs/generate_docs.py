"""
Documentation Generator for Aetheria Celestial Sky Observer.
Generates comprehensive, professional PDF and DOCX architecture guides.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

DOCS_DIR = os.path.dirname(os.path.abspath(__file__))
PDF_PATH = os.path.join(DOCS_DIR, "ASTRONOMY_APP_ARCHITECTURE_GUIDE.pdf")
DOCX_PATH = os.path.join(DOCS_DIR, "ASTRONOMY_APP_ARCHITECTURE_GUIDE.docx")

def build_docx():
    """Generates a beautifully formatted Microsoft Word (DOCX) manual."""
    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Helvetica'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(40, 44, 52)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("AETHERIA CELESTIAL OBSERVER\n")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(18, 22, 31)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run("Full-Stack Architecture & Engineering Guide • Apple UI/UX Design System")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(255, 159, 10) # Apple Orange

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run("Author: Senior Full Stack Web Developer & Systems Architect • Production Release v2.0\n")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(120, 125, 135)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Table of Contents Summary
    doc.add_heading("1. Executive Summary & Design Vision", level=1)
    p = doc.add_paragraph(
        "Aetheria is an Apple-grade, high-performance celestial sky observation web application that renders "
        "a mathematically accurate, real-time 360-degree view of the cosmos from any vantage point on Earth. "
        "Strictly adhering to Apple's Human Interface Guidelines (HIG), the application fuses deep midnight aesthetics "
        "with translucent glassmorphism (Cupertino material blur), razor-sharp high-DPI instrumentation, and "
        "an organic, relaxing celestial soundscape generated in real time using the Web Audio API."
    )
    p.paragraph_format.line_spacing = 1.25

    p2 = doc.add_paragraph(
        "Unlike casual sky maps that render randomized decorative dots, Aetheria is anchored in verified astronomical "
        "astrometry sourced from the Yale Bright Star Catalog (BSC5) and Hipparcos J2000 epoch data. Every star, planet, "
        "constellation segment, and lunar phase is mathematically resolved against observer coordinates (Latitude and Longitude) "
        "and temporal epoch (Local Sidereal Time)."
    )
    p2.paragraph_format.line_spacing = 1.25

    # System Architecture
    doc.add_heading("2. Full-Stack System Architecture", level=1)
    p_arch = doc.add_paragraph(
        "The application employs a decoupled, low-latency client-server topology engineered for 60 FPS rendering "
        "and offline fault tolerance:\n"
        "• Frontend Presentation Tier: High-DPI HTML5 Canvas 2D/WebGL rendering pipeline, responsive Apple glassmorphic UI, "
        "linear Azimuth & Altitude tape scales, QuickTime-style timeline scrubber, and Web Audio API synthesizer.\n"
        "• Backend Computation Tier: FastAPI high-concurrency Python service calculating Keplerian planetary orbits, "
        "Julian dates, lunar phases, and catalog serialization.\n"
        "• Client-Side Astronomical Redundancy: Full autonomous calculation capability embedded client-side in astronomy.js "
        "so orientation panning, pitch scrubbing, and real-time twinkling operate with 0ms network latency."
    )
    p_arch.paragraph_format.line_spacing = 1.25

    # Architecture Table
    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Shading Accent 1'

    headers = ["Layer", "Technologies", "Key Responsibilities"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True

    rows_data = [
        ("Presentation", "HTML5, CSS3, SF Pro, Canvas 2D", "360° Spherical projection, High-DPI tape compasses, Glass UI"),
        ("Audio Engine", "Web Audio API, BiquadFilter, LFO", "432Hz ambient drone, resonant breathing filter, stardust chimes"),
        ("Time Engine", "Pointer Events API, UTC epoch", "Apple play bar, glued thumb scrubber, Cupertino date picker"),
        ("Backend Services", "FastAPI, Uvicorn, Python 3", "Astronomical ephemerides, star catalog API, PDF/DOCX downloads")
    ]

    for row_idx, data in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(data):
            table.cell(row_idx, col_idx).text = text

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Celestial Mathematics
    doc.add_heading("3. Celestial Mechanics & Astronomical Mathematics", level=1)
    doc.add_paragraph(
        "The accuracy of the sky view relies on rigorous spherical trigonometry and orbital mechanics equations:"
    )

    doc.add_heading("3.1 Julian Date Conversion", level=2)
    doc.add_paragraph(
        "For any UTC Gregorian year Y, month M, and fractional day D (including hours, minutes, seconds):\n"
        "If M <= 2, set Y = Y - 1 and M = M + 12.\n"
        "A = floor(Y / 100), B = 2 - A + floor(A / 4)\n"
        "JD = floor(365.25 * (Y + 4716)) + floor(30.6001 * (M + 1)) + D + B - 1524.5"
    )

    doc.add_heading("3.2 Sidereal Time & Local Hour Angle", level=2)
    doc.add_paragraph(
        "Greenwich Mean Sidereal Time (GMST) at J2000 epoch (T = (JD - 2451545.0) / 36525):\n"
        "GMST = 280.46061837 + 360.98564736629 * (JD - 2451545.0) + 0.000387933 * T² - T³ / 38710000.0\n"
        "Local Sidereal Time (LST) = GMST + Longitude (East positive)\n"
        "Hour Angle (H) = LST - Right Ascension (RA)"
    )

    doc.add_heading("3.3 Equatorial to Horizontal Coordinates (Alt/Az)", level=2)
    doc.add_paragraph(
        "Given observer latitude φ, object declination δ, and hour angle H:\n"
        "sin(h) = sin(δ) · sin(φ) + cos(δ) · cos(φ) · cos(H)\n"
        "h = arcsin(sin(h))\n"
        "cos(A) = [sin(δ) - sin(φ) · sin(h)] / [cos(φ) · cos(h)]\n"
        "If sin(H) > 0, then Azimuth A = 360° - arccos(cos(A)), else A = arccos(cos(A))."
    )

    doc.add_heading("3.4 3D Spherical Perspective Projection onto Screen Viewport", level=2)
    doc.add_paragraph(
        "Let camera orientation be (Acam, hcam) and Field of View be FOV. For an object at (A, h):\n"
        "ΔA = A - Acam\n"
        "vx = cos(h) · sin(ΔA)\n"
        "vy = sin(h) · cos(hcam) - cos(h) · sin(hcam) · cos(ΔA)\n"
        "vz = sin(h) · sin(hcam) + cos(h) · cos(hcam) · cos(ΔA)\n"
        "When vz > 0 (object in front of camera lens):\n"
        "x = ScreenWidth / 2 + (vx / vz) · [ScreenWidth / (2 · tan(FOV / 2))]\n"
        "y = ScreenHeight / 2 - (vy / vz) · [ScreenWidth / (2 · tan(FOV / 2))]"
    )

    # Frontend Engineering & UI
    doc.add_heading("4. Frontend Engineering & Apple Design System", level=1)
    doc.add_paragraph(
        "4.1 High-DPI Anti-Blur Canvas Rendering:\n"
        "Standard HTML canvas elements suffer severe blurriness on Apple Retina displays due to the mismatch "
        "between CSS pixels and hardware device pixels. Aetheria calculates window.devicePixelRatio on every resize event, "
        "multiplies canvas buffer dimensions (canvas.width = cssW * dpr), and normalizes the 2D transformation matrix "
        "(ctx.scale(dpr, dpr)). All instrument ticks and text are snapped with +0.5px sub-pixel alignment for razor sharpness."
    )

    doc.add_paragraph(
        "4.2 Linear Azimuth (Top) and Altitude (Left) Tape Compasses:\n"
        "• Top Linear Compass: Spans 360° horizontally with tick marks every 1° (fine), 5° (medium), and 10° (labeled). "
        "Cardinal headings (N, NE, E, SE, S, SW, W, NW) are rendered prominently, with North highlighted in Apple Orange (#ff9f0a). "
        "A live digital heading badge (e.g. '184° S') floats above the center indicator.\n"
        "• Left Linear Altitude Compass: Vertical tape spanning -90° (Nadir) to +90° (Zenith), with a vibrant green marker "
        "at 0° HORIZON and a digital elevation readout."
    )

    doc.add_paragraph(
        "4.3 Timeline Scrubber Stickiness (Zero-Lag Pointer Physics):\n"
        "To satisfy the strict user requirement ('when I slide the bar the circle moves first, make sure the circle sticks with the bar'), "
        "the scrubber utilizes Pointer Events API (pointerdown, pointermove, pointerup) combined with element.setPointerCapture(pointerId). "
        "CSS transitions are explicitly disabled on the thumb during dragging, locking its center directly to (clientX - rect.left) / rect.width. "
        "Calculations for hour, minute, second and celestial coordinate recalculation occur synchronously in the same frame."
    )

    doc.add_paragraph(
        "4.4 Apple iOS & macOS Cupertino 3D Drumroll Date Picker with Haptics Feedback:\n"
        "To provide authentic Apple interaction for celestial epoch selection, the application features a custom 3D Cupertino drumroll wheel picker:\n"
        "• 3D Cylindrical Perspective: Columns for Month, Day, Year, Hour, and Minute render along a 3D cylindrical arc with rotateX, Z-translation, and optical foreshortening.\n"
        "• Scroll-To-Rewind Interaction: Rolling or scrolling down on the drumroll or directly over the play bar date readout moves earlier dates into the central selection lens, scrubbing backward in time.\n"
        "• Dual-Tier Haptic Feedback: Combines physical device vibration via navigator.vibrate(8) for touch devices with a synthesized 12ms acoustic mechanical notch click detent via the Web Audio API on desktop and macOS.\n"
        "• Dynamic Calendar Constraints: Automatically reconciles days per month and leap years (clamping February to 28 or 29 days dynamically).\n"
        "• Quick Astronomical Presets: One-tap jumps to celestial highlights (Current Real Time, Midnight Tonight, Perseid Meteor Peak, Summer & Winter Solstices)."
    )

    doc.add_paragraph(
        "4.5 Apple Iconography & Controls:\n"
        "• White Gemini Logo: Elegant 4-pointed celestial spark / twin glyph representing constellations.\n"
        "• Flat White Minimal Telescope Logo: Clean, modern silhouette activating 3.5× magnification and a circular mil-dot reticle.\n"
        "• Asteroid / Meteorite Logo: Stylized faceted meteor glyph for minor planets and comets.\n"
        "• Minimal Apple Sound Icon: SF Symbols style speaker with live animated sound waves pulsating when ambience is active.\n"
        "• Vertical Right Action Dock: Borderless 30% translucent liquid glass buttons on the right edge with zero 3D bulge."
    )

    doc.add_paragraph(
        "4.6 Ambient Relaxing Sound Synthesizer:\n"
        "Constructed with Web Audio API nodes: a 432Hz fundamental drone, dual detuned sine/triangle oscillators for binaural warmth, "
        "a 380Hz Moog-style biquad lowpass filter with slow 0.065Hz LFO modulation, and randomized pentatonic stardust chimes (864Hz - 1728Hz) "
        "with stereo spatial panning. Smooth exponential gain ramps prevent any clicks or pops."
    )

    # Backend Engineering
    doc.add_heading("5. Backend Engineering & Microservices", level=1)
    doc.add_paragraph(
        "The backend is developed with FastAPI and Uvicorn, exposing high-speed RESTful endpoints:\n"
        "• GET /api/ephemeris: Resolves Sun, Moon, and 5 planets (Mercury through Saturn) plus bright stars for given timestamp and location.\n"
        "• GET /api/stars: Returns the immutable star and constellation geometry database.\n"
        "• GET /api/time: Returns synchronized atomic UTC time and Julian Date.\n"
        "• GET /api/download/pdf & /api/download/docx: Serves the architecture documentation."
    )

    # Step-by-Step Build Procedure
    doc.add_heading("6. Step-by-Step Build & Deployment Procedure", level=1)
    doc.add_paragraph(
        "Step 1: Install Python prerequisites:\n"
        "   pip3 install fastapi uvicorn reportlab python-docx pillow\n\n"
        "Step 2: Initialize project directory structure:\n"
        "   backend/ (astronomy_engine.py, star_catalog.py, app.py)\n"
        "   frontend/ (index.html, css/apple-theme.css, css/compass.css, js/)\n"
        "   docs/ (generate_docs.py, ASTRONOMY_APP_ARCHITECTURE_GUIDE.pdf/docx)\n\n"
        "Step 3: Launch FastAPI server:\n"
        "   python3 -m uvicorn backend.app:app --host 0.0.0.0 --port 8000 --reload\n\n"
        "Step 4: Open web browser to http://localhost:8000\n"
        "   Interact with 360° celestial sphere, linear compasses, audio synthesizer, and timeline scrubber."
    )

    doc.save(DOCX_PATH)
    print(f"Successfully generated DOCX at: {DOCX_PATH}")

def build_pdf():
    """Generates a professional, high-craft PDF document using ReportLab."""
    pdf = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#12161f'),
        alignment=1,
        spaceAfter=6
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#ff9f0a'), # Apple orange
        alignment=1,
        spaceAfter=4
    )

    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#707582'),
        alignment=1,
        spaceAfter=18
    )

    h1_style = ParagraphStyle(
        'SectionH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#12161f'),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'SectionH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#ff9f0a'),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#2c3038'),
        spaceAfter=8
    )

    code_style = ParagraphStyle(
        'CodeSnippet',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#1a202c'),
        backColor=colors.HexColor('#f7f8fa'),
        borderColor=colors.HexColor('#e2e8f0'),
        borderWidth=0.5,
        borderPadding=6,
        spaceBefore=4,
        spaceAfter=8
    )

    elements = []

    # Title & Subtitle
    elements.append(Paragraph("AETHERIA CELESTIAL OBSERVER", title_style))
    elements.append(Paragraph("Full-Stack Architecture & Engineering Guide • Apple UI/UX Design System", subtitle_style))
    elements.append(Paragraph("Production Grade Architecture • Senior Full Stack Web Developer & Systems Engineer", meta_style))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e2e8f0'), spaceAfter=14))

    # 1. Executive Summary
    elements.append(Paragraph("1. Executive Summary & Design Vision", h1_style))
    elements.append(Paragraph(
        "Aetheria is an Apple-grade, high-performance celestial sky observation web application that renders a "
        "mathematically accurate, real-time 360-degree view of the cosmos from any vantage point on Earth. "
        "Strictly adhering to Apple's Human Interface Guidelines (HIG), the application fuses deep midnight aesthetics "
        "with translucent glassmorphism (Cupertino material blur), razor-sharp high-DPI instrumentation, and "
        "an organic, relaxing celestial soundscape generated in real time using the Web Audio API.",
        body_style
    ))
    elements.append(Paragraph(
        "Unlike casual sky apps that display randomized decorative dots, Aetheria is anchored in verified astronomical "
        "astrometry sourced from the Yale Bright Star Catalog (BSC5) and Hipparcos J2000 epoch data. Every star, planet, "
        "constellation segment, and lunar phase is mathematically resolved against observer coordinates (Latitude and Longitude) "
        "and temporal epoch (Local Sidereal Time).",
        body_style
    ))

    # 2. System Architecture
    elements.append(Paragraph("2. Full-Stack System Architecture", h1_style))
    elements.append(Paragraph(
        "The application employs a decoupled, low-latency client-server topology engineered for 60 FPS rendering "
        "and offline fault tolerance:",
        body_style
    ))

    # Architecture Table
    data = [
        [Paragraph("<b>Tier</b>", body_style), Paragraph("<b>Components & Tech</b>", body_style), Paragraph("<b>Key Responsibilities</b>", body_style)],
        [Paragraph("<b>Frontend Presentation</b>", body_style), Paragraph("HTML5 Canvas 2D, CSS3, SF Pro, Vanilla JS", body_style), Paragraph("360° Spherical projection, High-DPI tape compasses, Glass UI", body_style)],
        [Paragraph("<b>Audio Engine</b>", body_style), Paragraph("Web Audio API, BiquadFilter, LFO", body_style), Paragraph("432Hz ambient drone, resonant breathing filter, stardust chimes", body_style)],
        [Paragraph("<b>Time Engine</b>", body_style), Paragraph("Pointer Events API, UTC epoch math", body_style), Paragraph("Apple play bar, glued thumb scrubber, Cupertino date picker", body_style)],
        [Paragraph("<b>Backend Services</b>", body_style), Paragraph("FastAPI, Uvicorn, Python 3.14", body_style), Paragraph("Astronomical ephemerides, star catalog API, document delivery", body_style)]
    ]

    table = Table(data, colWidths=[120, 160, 220])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#f1f3f8')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#12161f')),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 10))

    # 3. Celestial Mathematics
    elements.append(Paragraph("3. Celestial Mechanics & Astronomical Mathematics", h1_style))
    elements.append(Paragraph("3.1 Julian Date Conversion", h2_style))
    elements.append(Paragraph(
        "For any Gregorian UTC timestamp (Y, M, D, H, Min, Sec):<br/>"
        "If M &le; 2: Y = Y - 1, M = M + 12<br/>"
        "A = floor(Y / 100), B = 2 - A + floor(A / 4)<br/>"
        "JD = floor(365.25 &times; (Y + 4716)) + floor(30.6001 &times; (M + 1)) + D + B - 1524.5",
        code_style
    ))

    elements.append(Paragraph("3.2 Sidereal Time & Local Hour Angle", h2_style))
    elements.append(Paragraph(
        "GMST = 280.46061837 + 360.98564736629 &times; (JD - 2451545.0) + 0.000387933 &times; T&sup2; - T&sup3; / 38710000.0<br/>"
        "LST = (GMST + Longitude) mod 360&deg;<br/>"
        "Hour Angle H = (LST - Right Ascension) mod 360&deg;",
        code_style
    ))

    elements.append(Paragraph("3.3 Equatorial to Horizontal Coordinates (Alt/Az)", h2_style))
    elements.append(Paragraph(
        "sin(h) = sin(&delta;) &times; sin(&phi;) + cos(&delta;) &times; cos(&phi;) &times; cos(H)<br/>"
        "Altitude h = arcsin(sin(h))<br/>"
        "cos(Az) = [sin(&delta;) - sin(&phi;) &times; sin(h)] / [cos(&phi;) &times; cos(h)]<br/>"
        "If sin(H) &gt; 0: Azimuth = 360&deg; - arccos(cos(Az)), else Azimuth = arccos(cos(Az))",
        code_style
    ))

    elements.append(Paragraph("3.4 3D Spherical Perspective Projection onto Screen Viewport", h2_style))
    elements.append(Paragraph(
        "Let camera orientation be (Acam, hcam) and Field of View be FOV. For an object at (A, h):<br/>"
        "&Delta;A = A - Acam<br/>"
        "vx = cos(h) &times; sin(&Delta;A)<br/>"
        "vy = sin(h) &times; cos(hcam) - cos(h) &times; sin(hcam) &times; cos(&Delta;A)<br/>"
        "vz = sin(h) &times; sin(hcam) + cos(h) &times; cos(hcam) &times; cos(&Delta;A)<br/>"
        "When vz &gt; 0 (object in front of camera):<br/>"
        "x = ScreenWidth / 2 + (vx / vz) &times; [ScreenWidth / (2 &times; tan(FOV / 2))]<br/>"
        "y = ScreenHeight / 2 - (vy / vz) &times; [ScreenWidth / (2 &times; tan(FOV / 2))]",
        code_style
    ))

    # 4. Frontend Engineering & Apple UI
    elements.append(Paragraph("4. Frontend Engineering & Apple Design System", h1_style))
    elements.append(Paragraph(
        "<b>High-DPI Anti-Blur Canvas Rendering:</b> Standard HTML canvas elements suffer severe blurriness on Apple Retina displays due to the mismatch between CSS pixels and hardware device pixels. Aetheria calculates window.devicePixelRatio on every resize event, scales canvas buffer dimensions (canvas.width = cssW * dpr), and normalizes the 2D transformation matrix (ctx.scale(dpr, dpr)). All instrument ticks and text are snapped with +0.5px sub-pixel alignment for razor sharpness.",
        body_style
    ))

    elements.append(Paragraph(
        "<b>Linear Azimuth (Top) and Altitude (Left) Tape Compasses:</b><br/>"
        "• Top Linear Compass: Spans 360° horizontally with tick marks every 1° (fine), 5° (medium), and 10° (labeled). Cardinal headings (N, NE, E, SE, S, SW, W, NW) are rendered prominently, with North highlighted in Apple Orange (#ff9f0a). A live digital heading badge (e.g. '184° S') floats above the center indicator.<br/>"
        "• Left Linear Altitude Compass: Vertical tape spanning -90° (Nadir) to +90° (Zenith), with a vibrant green marker at 0° HORIZON and a digital elevation readout.",
        body_style
    ))

    elements.append(Paragraph(
        "<b>Timeline Scrubber Stickiness (Zero-Lag Pointer Physics):</b> To satisfy the strict user requirement ('when I slide the bar the circle moves first, make sure the circle sticks with the bar'), the scrubber utilizes Pointer Events API (pointerdown, pointermove, pointerup) combined with element.setPointerCapture(pointerId). CSS transitions are explicitly disabled on the thumb during dragging, locking its center directly to (clientX - rect.left) / rect.width. Calculations for hour, minute, second and celestial coordinate recalculation occur synchronously in the same frame.",
        body_style
    ))

    elements.append(Paragraph(
        "<b>Apple iOS & macOS Cupertino 3D Drumroll Date Picker with Haptics Feedback:</b><br/>"
        "• 3D Cylindrical Perspective: Columns for Month, Day, Year, Hour, and Minute render along a 3D cylindrical arc with rotateX, Z-translation, and optical foreshortening.<br/>"
        "• Scroll-To-Rewind Interaction: Rolling or scrolling down on the drumroll or directly over the play bar date readout moves earlier dates into the central selection lens, scrubbing backward in time.<br/>"
        "• Dual-Tier Haptic Feedback: Combines physical device vibration via navigator.vibrate(8) for touch devices with a synthesized 12ms acoustic mechanical notch click detent via the Web Audio API on desktop and macOS.<br/>"
        "• Dynamic Calendar Constraints: Automatically reconciles days per month and leap years (clamping February to 28 or 29 days dynamically).<br/>"
        "• Quick Astronomical Presets: One-tap jumps to celestial highlights (Current Real Time, Midnight Tonight, Perseid Meteor Peak, Summer & Winter Solstices).",
        body_style
    ))

    elements.append(Paragraph(
        "<b>Apple Iconography & Controls:</b><br/>"
        "• White Gemini Logo: Elegant 4-pointed celestial spark / twin glyph representing constellations.<br/>"
        "• Flat White Minimal Telescope Logo: Clean, modern silhouette activating 3.5× magnification and a circular mil-dot reticle.<br/>"
        "• Asteroid / Meteorite Logo: Stylized faceted meteor glyph for minor planets and comets.<br/>"
        "• Minimal Apple Sound Icon: SF Symbols style speaker with live animated sound waves pulsating when ambience is active.<br/>"
        "• Vertical Right Action Dock: Borderless 30% translucent liquid glass buttons on the right edge with zero 3D bulge.",
        body_style
    ))

    elements.append(Paragraph(
        "<b>Ambient Relaxing Sound Synthesizer:</b> Constructed with Web Audio API nodes: a 432Hz fundamental drone, dual detuned sine/triangle oscillators for binaural warmth, a 380Hz Moog-style biquad lowpass filter with slow 0.065Hz LFO modulation, and randomized pentatonic stardust chimes (864Hz - 1728Hz) with stereo spatial panning. Smooth exponential gain ramps prevent any clicks or pops.",
        body_style
    ))

    # 5. Backend Engineering
    elements.append(Paragraph("5. Backend Engineering & Microservices", h1_style))
    elements.append(Paragraph(
        "The backend is developed with FastAPI and Uvicorn, exposing high-speed RESTful endpoints:<br/>"
        "• GET /api/ephemeris: Resolves Sun, Moon, and 5 planets (Mercury through Saturn) plus bright stars for given timestamp and location.<br/>"
        "• GET /api/stars: Returns the immutable star and constellation geometry database.<br/>"
        "• GET /api/time: Returns synchronized atomic UTC time and Julian Date.<br/>"
        "• GET /api/download/pdf & /api/download/docx: Serves the architecture documentation.",
        body_style
    ))

    # 6. Step-by-Step Procedure
    elements.append(Paragraph("6. Step-by-Step Build & Deployment Procedure", h1_style))
    elements.append(Paragraph(
        "Step 1: Install Python prerequisites:<br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;pip3 install fastapi uvicorn reportlab python-docx pillow<br/><br/>"
        "Step 2: Initialize project directory structure:<br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;backend/ (astronomy_engine.py, star_catalog.py, app.py)<br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;frontend/ (index.html, css/apple-theme.css, css/compass.css, js/)<br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;docs/ (generate_docs.py, ASTRONOMY_APP_ARCHITECTURE_GUIDE.pdf/docx)<br/><br/>"
        "Step 3: Launch FastAPI server:<br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;python3 -m uvicorn backend.app:app --host 0.0.0.0 --port 8000 --reload<br/><br/>"
        "Step 4: Open web browser to http://localhost:8000<br/>"
        "&nbsp;&nbsp;&nbsp;&nbsp;Interact with 360° celestial sphere, linear compasses, audio synthesizer, and timeline scrubber.",
        code_style
    ))

    pdf.build(elements)
    print(f"Successfully generated PDF at: {PDF_PATH}")

if __name__ == "__main__":
    os.makedirs(DOCS_DIR, exist_ok=True)
    build_docx()
    build_pdf()
